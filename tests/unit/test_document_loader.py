"""services/document_loader.py — chuyển docx/pdf/txt/md sang Markdown."""

import io

import pytest
from docx import Document

from slidegen_rag.services.document_loader import (
    EmptyDocumentError,
    UnsupportedDocumentError,
    to_markdown,
)


def _minimal_pdf(pages: list[str]) -> bytes:
    """PDF tối giản, mỗi trang một dòng chữ Helvetica (không cần lib tạo PDF)."""
    objs: list[bytes] = []
    n = len(pages)
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(n))
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode())
    font_id = 3 + 2 * n
    for i, text in enumerate(pages):
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode() if text else b""
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {4 + 2 * i} 0 R /Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()
        )
        objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = b"%PDF-1.4\n"
    offsets = []
    for i, o in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def _docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Chương 1", level=1)
    doc.add_paragraph("Đoạn mở đầu.")
    doc.add_heading("Mục 1.1", level=2)
    doc.add_paragraph("Ý một", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    for r, row in enumerate([["Tên", "Giá"], ["A", "10"]]):
        for c, val in enumerate(row):
            table.cell(r, c).text = val
    doc.add_paragraph("Kết luận.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_markdown_and_txt_passthrough():
    assert to_markdown("a.md", b"# Hi\n\nx") == "# Hi\n\nx"
    assert to_markdown("a.txt", "xin chào".encode()) == "xin chào"


def test_docx_headings_lists_tables_in_order():
    md = to_markdown("a.docx", _docx_bytes())
    assert md.index("# Chương 1") < md.index("Đoạn mở đầu.") < md.index("## Mục 1.1")
    assert "- Ý một" in md
    assert "| Tên | Giá |" in md and "| --- | --- |" in md and "| A | 10 |" in md
    assert md.index("| A | 10 |") < md.index("Kết luận.")


def test_pdf_extracts_text_per_page():
    md = to_markdown("Báo cáo.pdf", _minimal_pdf(["Hello page one", "Second page"]))
    assert md.startswith("# Báo cáo")
    assert "## Trang 1" in md and "Hello page one" in md
    assert "## Trang 2" in md and "Second page" in md


def test_pdf_without_text_raises_empty():
    with pytest.raises(EmptyDocumentError):
        to_markdown("scan.pdf", _minimal_pdf([""]))


def test_unsupported_extension_raises():
    with pytest.raises(UnsupportedDocumentError):
        to_markdown("a.xlsx", b"x")


def test_empty_markdown_raises():
    with pytest.raises(EmptyDocumentError):
        to_markdown("a.md", b"   \n")
