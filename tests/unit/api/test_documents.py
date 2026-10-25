def _create_project(client, name="Demo"):
    return client.post("/api/v1/projects", json={"name": name}).json()["id"]


def _upload(client, project_id, title="Doc", content=b"# Hello\n\nWorld"):
    return client.post(
        f"/api/v1/projects/{project_id}/documents",
        data={"title": title},
        files={"file": ("doc.md", content, "text/markdown")},
    )


def test_upload_document(client):
    project_id = _create_project(client)
    resp = _upload(client, project_id)
    assert resp.status_code == 201
    body = resp.json()
    assert body["title"] == "Doc"
    assert body["project_id"] == project_id
    assert body["status"] == "queued"


def test_upload_dedupes_by_sha1(client):
    project_id = _create_project(client)
    first = _upload(client, project_id).json()
    second = _upload(client, project_id).json()
    assert first["id"] == second["id"]


def test_list_documents_in_project(client):
    project_id = _create_project(client)
    _upload(client, project_id, title="Doc A")
    resp = client.get(f"/api/v1/projects/{project_id}/documents")
    assert resp.status_code == 200
    assert len(resp.json()) == 1


def test_get_document_detail(client):
    project_id = _create_project(client)
    document_id = _upload(client, project_id).json()["id"]
    resp = client.get(f"/api/v1/documents/{document_id}")
    assert resp.status_code == 200
    body = resp.json()
    assert body["document"]["id"] == document_id
    assert body["active_index"] is None


def test_rename_and_delete_document(client):
    project_id = _create_project(client)
    document_id = _upload(client, project_id).json()["id"]

    resp = client.patch(f"/api/v1/documents/{document_id}", json={"title": "Renamed"})
    assert resp.status_code == 200
    assert resp.json()["title"] == "Renamed"

    resp = client.delete(f"/api/v1/documents/{document_id}")
    assert resp.status_code == 204

    # Soft delete: get_document (như bản Postgres) không lọc deleted_at, chỉ
    # list_documents mới ẩn document đã xoá — xem memory/catalog.py.
    resp = client.get(f"/api/v1/documents/{document_id}")
    assert resp.status_code == 200
    assert resp.json()["document"]["deleted_at"] is not None

    resp = client.get(f"/api/v1/projects/{project_id}/documents")
    assert document_id not in {d["id"] for d in resp.json()}


def test_start_index_returns_job_id(client):
    project_id = _create_project(client)
    document_id = _upload(client, project_id).json()["id"]

    resp = client.post(f"/api/v1/documents/{document_id}/index")
    assert resp.status_code == 202
    body = resp.json()
    assert body["job_id"]
    assert body["index_id"]


def test_start_index_conflict_when_already_parsing(client):
    project_id = _create_project(client)
    document_id = _upload(client, project_id).json()["id"]

    first = client.post(f"/api/v1/documents/{document_id}/index")
    assert first.status_code == 202

    second = client.post(f"/api/v1/documents/{document_id}/index")
    assert second.status_code == 409


def test_start_index_missing_document_404(client):
    resp = client.post("/api/v1/documents/00000000-0000-0000-0000-000000000099/index")
    assert resp.status_code == 404


def test_upload_docx_is_converted_to_markdown(client):
    import io

    from docx import Document

    buf = io.BytesIO()
    doc = Document()
    doc.add_heading("Tiêu đề", level=1)
    doc.add_paragraph("Nội dung")
    doc.save(buf)
    project_id = client.post("/api/v1/projects", json={"name": "P"}).json()["id"]
    resp = client.post(
        f"/api/v1/projects/{project_id}/documents",
        data={"title": "Docx"},
        files={"file": ("a.docx", buf.getvalue(), "application/octet-stream")},
    )
    assert resp.status_code == 201


def test_upload_unsupported_extension_returns_422(client):
    project_id = client.post("/api/v1/projects", json={"name": "P"}).json()["id"]
    resp = client.post(
        f"/api/v1/projects/{project_id}/documents",
        data={"title": "X"},
        files={"file": ("a.xlsx", b"x", "application/octet-stream")},
    )
    assert resp.status_code == 422
